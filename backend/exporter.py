import os
import csv
import time
from datetime import datetime
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Define the exact 2-tier header structure matching to_follow.xlsx (28 columns)
EXPORT_COLUMNS_R1 = [
    "Sr", "NAAC ", None, None,
    "Strategic Plan ", None, None,
    "Gradguate Attribute", None, None,
    "General Information :", None, None, None, None, None,
    "Speaker/Guest/Presenter Details :", None, None, None,
    "Participant’s profile :", None,
    "Synopsis of the Activity (Description):", None, None, None,
    "Rapporteur Details:", None
]

EXPORT_COLUMNS_R2 = [
    None,
    "Criteria/ Focus Area\nStandard No.",
    "Sub Criteria/ \nStandard No",
    "Sub-Criteria/ \nStandard Title",
    "Criteria/ Focus Area\nStandard No.",
    "Sub Criteria/ \nStandard No",
    "Sub-Criteria/ \nStandard Title",
    "Criteria/ Focus Area\nStandard No.",
    "Sub Criteria/ \nStandard No",
    "Sub-Criteria/ \nStandard Title",
    "Type of Activity ",
    "Title of the Activity ",
    "Date/s ",
    "Time ",
    "Venue ",
    "    Collaboration/Sponsor  (if any) ",
    "Name ",
    "Title/Position ",
    "Organization ",
    "Title of Presentation ",
    "Type of Participants ",
    "No. of Participants",
    "Highlights of the Activity",
    "Key Objectives/Takeaways",
    "Summary of the Activity",
    "Follow-up Plan, if any ",
    "Name of the Rapporteur ",
    "Email and Contact No "
]

# Style definitions extracted directly from to_follow.xlsx
CELL_STYLES_R1 = [
    ('Aptos Narrow', 12.0, True, None, None, None),
    ('Aptos Narrow', 12.0, True, 'center', 'center', None),
    ('Aptos Narrow', 12.0, True, 'center', 'center', None),
    ('Aptos Narrow', 12.0, True, 'center', 'center', None),
    ('Georgia', 12.0, True, 'center', 'center', None),
    ('Aptos Narrow', 12.0, True, 'center', 'center', None),
    ('Aptos Narrow', 12.0, True, 'center', 'center', None),
    ('Georgia', 12.0, True, 'center', 'center', None),
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Georgia', 12.0, True, None, 'center', None),
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Georgia', 12.0, True, None, 'center', None),
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Georgia', 12.0, True, None, 'center', None),
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Georgia', 12.0, True, None, None, None),
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Georgia', 12.0, True, None, 'center', None),
    ('Aptos Narrow', 11.0, False, None, None, None),
]

CELL_STYLES_R2 = [
    ('Aptos Narrow', 11.0, False, None, None, None),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
    ('Georgia', 11.0, False, 'center', 'center', True),
]

COLUMN_WIDTHS = [
    13.0, 25.8, 21.6, 17.6, 20.4, 12.0, 34.4, 28.2, 11.8, 19.2,
    32.1, 39.6, 14.0, 31.3, 26.2, 30.0, 40.4, 51.0, 19.0, 16.1,
    25.6, 19.8, 44.7, 28.8, 28.0, 28.0, 13.0, 13.0
]

# Field extractors corresponding to columns 1 to 28
# Note: col 1 is 'Sr' (serial number)
COLUMN_EXTRACTORS = [
    lambda r, idx: str(idx),
    # NAAC
    lambda r, idx: r.get("academic_mapping", {}).get("naac", {}).get("criteria", ""),
    lambda r, idx: r.get("academic_mapping", {}).get("naac", {}).get("sub_criteria_no", ""),
    lambda r, idx: r.get("academic_mapping", {}).get("naac", {}).get("sub_criteria_title", ""),
    # Strategic Plan
    lambda r, idx: r.get("academic_mapping", {}).get("strategic_plan", {}).get("criteria", ""),
    lambda r, idx: r.get("academic_mapping", {}).get("strategic_plan", {}).get("sub_criteria_no", ""),
    lambda r, idx: r.get("academic_mapping", {}).get("strategic_plan", {}).get("sub_criteria_title", ""),
    # Graduate Attribute
    lambda r, idx: r.get("academic_mapping", {}).get("graduate_attribute", {}).get("criteria", ""),
    lambda r, idx: r.get("academic_mapping", {}).get("graduate_attribute", {}).get("sub_criteria_no", ""),
    lambda r, idx: r.get("academic_mapping", {}).get("graduate_attribute", {}).get("sub_criteria_title", ""),
    # General Information
    lambda r, idx: r.get("general_information", {}).get("type", ""),
    lambda r, idx: r.get("general_information", {}).get("title", ""),
    lambda r, idx: r.get("general_information", {}).get("date", ""),
    lambda r, idx: r.get("general_information", {}).get("time", ""),
    lambda r, idx: r.get("general_information", {}).get("venue", ""),
    lambda r, idx: r.get("general_information", {}).get("collaboration", ""),
    # Speaker
    lambda r, idx: r.get("speaker_details", {}).get("name", ""),
    lambda r, idx: r.get("speaker_details", {}).get("position", ""),
    lambda r, idx: r.get("speaker_details", {}).get("organization", ""),
    lambda r, idx: r.get("speaker_details", {}).get("presentation_title", ""),
    # Participants
    lambda r, idx: r.get("participant_profile", {}).get("participant_type", ""),
    lambda r, idx: r.get("participant_profile", {}).get("participant_count", ""),
    # Synopsis / Description
    lambda r, idx: r.get("highlights", ""),
    lambda r, idx: r.get("key_objectives", ""),
    lambda r, idx: r.get("summary", ""),
    lambda r, idx: r.get("follow_up_plan", "None"),
    # Rapporteur
    lambda r, idx: r.get("rapporteur_details", {}).get("name", ""),
    lambda r, idx: r.get("rapporteur_details", {}).get("contact", "")
]

# Combined 1-line display headers for simple reference / previews
COMBINED_HEADERS = [
    "Sr",
    "NAAC - Criteria/Standard No.",
    "NAAC - Sub Criteria No.",
    "NAAC - Sub-Criteria Title",
    "Strategic Plan - Criteria/Standard No.",
    "Strategic Plan - Sub Criteria No.",
    "Strategic Plan - Sub-Criteria Title",
    "Graduate Attribute - Criteria/Standard No.",
    "Graduate Attribute - Sub Criteria No.",
    "Graduate Attribute - Sub-Criteria Title",
    "Type of Activity",
    "Title of the Activity",
    "Date/s",
    "Time",
    "Venue",
    "Collaboration/Sponsor",
    "Speaker Name",
    "Speaker Title/Position",
    "Speaker Organization",
    "Speaker Presentation Title",
    "Type of Participants",
    "No. of Participants",
    "Highlights of the Activity",
    "Key Objectives/Takeaways",
    "Summary of the Activity",
    "Follow-up Plan",
    "Name of the Rapporteur",
    "Email and Contact No"
]

def record_to_row(record: Dict[str, Any], serial_number: int) -> List[str]:
    """Extracts row values strictly following the 28 columns of to_follow.xlsx."""
    row = []
    for fn in COLUMN_EXTRACTORS:
        val = str(fn(record, serial_number) or "").strip()
        row.append(val)
    return row

def check_file_collision(file_path: str) -> Dict[str, Any]:
    """
    Checks if target file already exists and returns its metadata for the conflict dialog.
    """
    if not os.path.exists(file_path):
        return {
            "exists": False,
            "file_name": os.path.basename(file_path),
            "file_path": os.path.abspath(file_path),
            "file_type": os.path.splitext(file_path)[1].lower()
        }
    
    stat = os.stat(file_path)
    file_type = os.path.splitext(file_path)[1].lower()
    
    row_count = 0
    if file_type == ".csv":
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                # 2 header rows in to_follow structure
                total_lines = sum(1 for _ in f)
                row_count = max(0, total_lines - 2)
        except Exception:
            row_count = -1
    elif file_type in [".xlsx", ".xls"]:
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True)
            ws = wb.active
            # Count rows starting from row 3 (after 2 header rows)
            data_rows = 0
            for row in ws.iter_rows(min_row=3, values_only=True):
                if any(row):
                    data_rows += 1
            row_count = data_rows
            wb.close()
        except Exception:
            row_count = -1

    return {
        "exists": True,
        "file_name": os.path.basename(file_path),
        "file_path": os.path.abspath(file_path),
        "file_type": file_type,
        "size_bytes": stat.st_size,
        "size_formatted": f"{stat.st_size / 1024:.1f} KB" if stat.st_size < 1024*1024 else f"{stat.st_size / (1024*1024):.2f} MB",
        "last_modified": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
        "existing_rows": row_count
    }

def get_unique_filename(file_path: str) -> str:
    """Generates a non-conflicting unique filename using timestamp."""
    base, ext = os.path.splitext(file_path)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{base}_{timestamp}{ext}"

def export_csv(file_path: str, records: List[Dict[str, Any]], mode: str = "new") -> str:
    """
    Exports records to CSV using the exact 28-column, 2-header-row structure of to_follow.xlsx.
    mode: 'append' to append to existing, 'new' to write to new unique file.
    """
    target_path = file_path
    if mode == "new" and os.path.exists(file_path):
        target_path = get_unique_filename(file_path)

    os.makedirs(os.path.dirname(os.path.abspath(target_path)), exist_ok=True)
    
    is_append = (mode == "append" and os.path.exists(target_path))
    start_sr = 1
    if is_append:
        # Ensure target file ends with newline so new row doesn't merge with last line
        try:
            with open(target_path, "rb+") as f:
                f.seek(0, os.SEEK_END)
                if f.tell() > 0:
                    f.seek(-1, os.SEEK_END)
                    if f.read(1) != b"\n":
                        f.write(b"\n")
        except Exception:
            pass

        try:
            with open(target_path, "r", encoding="utf-8", errors="ignore") as f:
                lines = [line for line in f if line.strip()]
                total_data_lines = max(0, len(lines) - 2)
                if total_data_lines > 0:
                    last_line = lines[-1]
                    first_cell = last_line.split(",")[0].strip().strip('"')
                    try:
                        start_sr = int(first_cell) + 1
                    except ValueError:
                        start_sr = total_data_lines + 1
                else:
                    start_sr = 1
        except Exception:
            start_sr = 1

    write_mode = "a" if is_append else "w"
    
    with open(target_path, write_mode, newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        if not is_append:
            # Write Row 1 (Section Groups)
            writer.writerow([c or "" for c in EXPORT_COLUMNS_R1])
            # Write Row 2 (Sub-headers)
            writer.writerow([c or "" for c in EXPORT_COLUMNS_R2])
        
        # Write Data Rows
        for i, r in enumerate(records):
            writer.writerow(record_to_row(r, start_sr + i))

    return target_path

def apply_to_follow_header_styles(ws):
    """Applies exact cell styling and column widths matching to_follow.xlsx."""
    # Row Heights
    ws.row_dimensions[1].height = 15.6
    ws.row_dimensions[2].height = 55.2

    # Column Widths
    for idx, width in enumerate(COLUMN_WIDTHS, start=1):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws.column_dimensions[col_letter].width = width

    # Row 1 headers and styles
    for idx, val in enumerate(EXPORT_COLUMNS_R1, start=1):
        cell = ws.cell(row=1, column=idx, value=val)
        font_name, font_size, bold, h_align, v_align, wrap = CELL_STYLES_R1[idx - 1]
        cell.font = Font(name=font_name, size=font_size, bold=bold)
        cell.alignment = Alignment(horizontal=h_align, vertical=v_align, wrap_text=wrap)

    # Row 2 headers and styles
    for idx, val in enumerate(EXPORT_COLUMNS_R2, start=1):
        cell = ws.cell(row=2, column=idx, value=val)
        font_name, font_size, bold, h_align, v_align, wrap = CELL_STYLES_R2[idx - 1]
        cell.font = Font(name=font_name, size=font_size, bold=bold)
        cell.alignment = Alignment(horizontal=h_align, vertical=v_align, wrap_text=wrap)

def export_excel(file_path: str, records: List[Dict[str, Any]], mode: str = "new") -> str:
    """
    Exports records to Excel (.xlsx) exactly structured and styled like to_follow.xlsx.
    mode: 'append' to append rows, 'new' to create fresh styled workbook.
    """
    target_path = file_path
    if mode == "new" and os.path.exists(file_path):
        target_path = get_unique_filename(file_path)

    os.makedirs(os.path.dirname(os.path.abspath(target_path)), exist_ok=True)

    if mode == "append" and os.path.exists(target_path):
        wb = openpyxl.load_workbook(target_path)
        ws = wb.active
        # Determine current max non-empty data row
        last_data_row = 2
        for r in range(3, ws.max_row + 1):
            if any(ws.cell(r, c).value for c in range(1, 29)):
                last_data_row = r
        
        # Determine starting serial number from previous row or count
        last_sr_val = ws.cell(last_data_row, 1).value
        try:
            start_sr = int(str(last_sr_val).strip()) + 1
        except (ValueError, TypeError):
            start_sr = max(1, last_data_row - 2 + 1)
        next_row = last_data_row + 1
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        apply_to_follow_header_styles(ws)
        start_sr = 1
        next_row = 3

    # Append data rows with clean styling
    data_font = Font(name="Georgia", size=10)
    for i, r in enumerate(records):
        row_vals = record_to_row(r, start_sr + i)
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=next_row, column=col_idx, value=val)
            cell.font = data_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        next_row += 1

    wb.save(target_path)
    return target_path
